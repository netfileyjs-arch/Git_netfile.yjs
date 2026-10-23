# -*- coding: utf-8 -*-
"""
파일명 키워드 필터 - 핵심 로직

- 키워드를 한글 ↔ 영문 사전(dictionary.txt)으로 서로 번역해서 확장
- 파일명 목록에서 키워드(원문 + 번역어)가 들어간 파일명을 골라냄
- 결과를 엑셀(xlsx)로 저장
"""

import io
import os
import re
import unicodedata
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DICTIONARY_PATH = os.path.join(BASE_DIR, 'dictionary.txt')

HANGUL_RE = re.compile(r'[가-힣ㄱ-ㆎᄀ-ᇿ]')
KEYWORD_SPLIT_RE = re.compile(r'[/\n\r,;|]+')
AGE_RE = re.compile(r'^(\d{1,2})\s*(세|살|歲|歳|yo|y\.o\.?|yearsold|years old|year old)$')
SEPARATOR_PATTERN = r'[\s._\-]*'


def normalize(text):
    """전각문자 → 보통문자(NFKC), 대소문자 무시"""
    return unicodedata.normalize('NFKC', text).casefold().strip()


def is_korean(term):
    return bool(HANGUL_RE.search(unicodedata.normalize('NFKC', term)))


# ──────────────────────────────────────────────
# 사전
# ──────────────────────────────────────────────

def parse_dictionary(text):
    """'한글1, 한글2 = eng1, eng2' 형식의 사전 텍스트 → [(ko_terms, en_terms), ...]"""
    pairs = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line.lstrip('='):
            continue
        left, right = _split_pair(line)
        ko = [t.strip() for t in left.split(',') if t.strip()]
        en = [t.strip() for t in right.split(',') if t.strip()]
        if ko and en:
            pairs.append((ko, en))
    return pairs


def _split_pair(line):
    """'=' 접두어(단어 단위 매칭)와 구분자 '=' 를 구별해서 나눔"""
    for match in re.finditer(r'=', line):
        i = match.start()
        before = line[:i].rstrip()
        # 구분자 '=' 는 앞에 한글/영문 단어가 있고, 쉼표 바로 뒤가 아님
        if before and not before.endswith(','):
            return line[:i], line[i + 1:]
    return line, ''


def load_dictionary(path=DICTIONARY_PATH):
    if not os.path.exists(path):
        return []
    with open(path, encoding='utf-8') as f:
        return parse_dictionary(f.read())


def build_index(pairs):
    """정규화된 단어 → 반대 언어 번역어 목록"""
    index = {}
    for ko_terms, en_terms in pairs:
        for term in ko_terms:
            index.setdefault(normalize(term.lstrip('=')), []).extend(en_terms)
        for term in en_terms:
            index.setdefault(normalize(term.lstrip('=')), []).extend(ko_terms)
    return index


# ──────────────────────────────────────────────
# 키워드
# ──────────────────────────────────────────────

def parse_keywords(text):
    """'/', 줄바꿈, 쉼표로 구분된 키워드 → 중복 제거된 목록 (입력 순서 유지)"""
    seen = set()
    keywords = []
    for raw in KEYWORD_SPLIT_RE.split(text):
        kw = raw.strip()
        key = normalize(kw)
        if kw and key not in seen:
            seen.add(key)
            keywords.append(kw)
    return keywords


def age_translations(keyword):
    """10세 / 13살 / 15歲 / 18yo / 18yearsold → 나머지 표기들"""
    m = AGE_RE.match(normalize(keyword))
    if not m:
        return []
    n = m.group(1)
    return [f'{n}세', f'{n}살', f'{n}歲', f'{n}歳',
            f'={n} yo', f'{n} years old', f'{n} year old', f'{n} yrs old']


def expand_keywords(keywords, pairs, translate=True):
    """각 키워드 → [(검색어, 구분)]  구분: '원문' 또는 '번역'"""
    index = build_index(pairs)
    expanded = {}
    for kw in keywords:
        terms = [(kw, '원문')]
        if translate:
            seen = {normalize(kw.lstrip('='))}
            candidates = index.get(normalize(kw.lstrip('=')), []) + age_translations(kw)
            for t in candidates:
                key = normalize(t.lstrip('='))
                if key not in seen:
                    seen.add(key)
                    terms.append((t, '번역'))
        expanded[kw] = terms
    return expanded


def compile_term(term):
    """검색어 → 정규식.  '=' 접두어는 영문 단어 단위 매칭, 띄어쓰기는 구분자 유무 무시"""
    whole_word = term.startswith('=') and len(term) > 1
    body = normalize(term[1:] if whole_word else term)
    words = [re.escape(w) for w in body.split()]
    pattern = SEPARATOR_PATTERN.join(words) if words else re.escape(body)
    if whole_word:
        pattern = r'(?<![a-z])' + pattern + r'(?![a-z])'
    return re.compile(pattern)


# ──────────────────────────────────────────────
# 필터링
# ──────────────────────────────────────────────

def parse_filenames(text):
    return [line.strip() for line in text.splitlines() if line.strip()]


def filter_filenames(filenames, keywords, pairs=None, translate=True):
    """
    반환: dict
      matched   : [{'no', 'name', 'keywords', 'terms', 'kinds'}]
      unmatched : [{'no', 'name'}]
      keywords  : [{'keyword', 'translations', 'count'}]
    """
    if pairs is None:
        pairs = load_dictionary()
    expanded = expand_keywords(keywords, pairs, translate)
    compiled = {kw: [(term, kind, compile_term(term)) for term, kind in terms]
                for kw, terms in expanded.items()}

    matched, unmatched = [], []
    counts = {kw: 0 for kw in keywords}
    for no, name in enumerate(filenames, start=1):
        target = normalize(name)
        hit_keywords, hit_terms, kinds = [], [], set()
        for kw, terms in compiled.items():
            for term, kind, regex in terms:
                if regex.search(target):
                    hit_keywords.append(kw)
                    hit_terms.append(term.lstrip('=') if kind == '번역' else term)
                    kinds.add(kind)
                    counts[kw] += 1
                    break
        if hit_keywords:
            matched.append({
                'no': no,
                'name': name,
                'keywords': hit_keywords,
                'terms': hit_terms,
                'kinds': '원문+번역' if len(kinds) == 2 else kinds.pop(),
            })
        else:
            unmatched.append({'no': no, 'name': name})

    keyword_rows = [{
        'keyword': kw,
        'translations': [t.lstrip('=') for t, kind in expanded[kw] if kind == '번역'],
        'count': counts[kw],
    } for kw in keywords]
    return {'matched': matched, 'unmatched': unmatched, 'keywords': keyword_rows}


# ──────────────────────────────────────────────
# 엑셀
# ──────────────────────────────────────────────

def build_workbook(result):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='305496')

    def write_sheet(ws, headers, rows, widths):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        for row in rows:
            ws.append(row)
        for col, width in zip('ABCDEFG', widths):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws = wb.active
    ws.title = '키워드 포함 (제거 대상)'
    write_sheet(ws, ['No', '원래 순번', '파일명', '매칭 키워드', '실제 걸린 단어', '구분'],
                [[i, m['no'], m['name'], ' / '.join(m['keywords']), ' / '.join(m['terms']), m['kinds']]
                 for i, m in enumerate(result['matched'], start=1)],
                [6, 10, 100, 30, 30, 10])

    ws = wb.create_sheet('나머지 파일')
    write_sheet(ws, ['No', '원래 순번', '파일명'],
                [[i, u['no'], u['name']] for i, u in enumerate(result['unmatched'], start=1)],
                [6, 10, 100])

    ws = wb.create_sheet('키워드 번역표')
    write_sheet(ws, ['키워드', '함께 검색한 번역어', '걸린 파일 수'],
                [[k['keyword'], ', '.join(k['translations']), k['count']] for k in result['keywords']],
                [24, 80, 12])

    ws = wb.create_sheet('요약')
    total = len(result['matched']) + len(result['unmatched'])
    for row in [['작성일시', datetime.now().strftime('%Y-%m-%d %H:%M')],
                ['전체 파일 수', total],
                ['키워드 포함 (제거 대상)', len(result['matched'])],
                ['나머지 파일', len(result['unmatched'])],
                ['키워드 수', len(result['keywords'])],
                ['한 개 이상 걸린 키워드 수', sum(1 for k in result['keywords'] if k['count'])]]:
        ws.append(row)
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 20
    return wb


def workbook_bytes(result):
    buffer = io.BytesIO()
    build_workbook(result).save(buffer)
    return buffer.getvalue()
